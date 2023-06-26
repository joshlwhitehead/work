"""This program reads temperature data collected from heatspreader tests, analyzes it, and makes a call to 'pass' or 'fail' the heatspreader.
PCR heat spreaders fail if the temperatures are less than 95% of the nominal temps.
TC heat spreaders fail if the temperatures are less than 80% of the nominal temps.
The program looks specifically at 80c, 90c, and 100c. An excel file is created which holds the data, calculations, plots, and pass/fail results.
'tkinter' is used to create a user interface so this code can be turned into a standalone exe"""

from tkinter import *                                                       
import os
import numpy as np
import pandas as pd
import openpyxl as op
from scipy.interpolate import interp1d
from openpyxl.chart import LineChart, Reference


##########################              SAVE DATA IN EXCEL              ########################
def toExcel(folder,longest,sameLenTime,sameLenTemp,newFile,tempc,timeTo2):                          # add formatted data to excel spreadsheet and graph it

    listIndx = np.arange(0,len(os.listdir(folder)))                                                 # list of numbers                   

    fullDict = {}                                                                                   #create dictionary to hold all data
    for fileIndx in range(len(os.listdir(folder))):                                                 #loop through files in data folder
        x = [os.listdir(folder)[fileIndx]]
    
        while len(x) < longest:                                                                     #make sure all data have same shape
            x.append(None)

        fullDict[''.join(['file name ',str(listIndx[fileIndx])])] = x                                   #add file name to dict
        fullDict[''.join(['normalized time (sec) ',str(listIndx[fileIndx])])] = sameLenTime[fileIndx]   #add time to dict
        fullDict[''.join(['temp (c) ',str(listIndx[fileIndx])])] = sameLenTemp[fileIndx]                #add temp to dict


    dFTot = pd.DataFrame(fullDict)                                              #create dataframe from dictionary
    writer = pd.ExcelWriter(newFile,engine='xlsxwriter')                        #use xlsxwriter to write data to excel
    dFTot.to_excel(writer,sheet_name='full')
    wb = writer.book
    ws = writer.sheets['full']
    chart = wb.add_chart({'type':'line'})                                       #add line chart to "full" sheet

    count = 2
    for tempIndx in range(len(sameLenTemp)):                                           #add data from "full" sheet to create line chart
        chart.add_series({
            'categories':['full',1,count,len(sameLenTemp[0]),count],
            'values':['full',1,count+1,len(sameLenTemp[0]),count+1],
            'name':['full',1,count-1]
            })
        count += 3
    
    chart.set_x_axis({'name':'Time (sec)'})                                 #label axes
    chart.set_y_axis({'name':'Temp (c)'})
    
    ws.insert_chart('D2',chart)
    writer.save()                                                           #save

    wb = op.load_workbook(newFile)                                          #load
    ws = wb.create_sheet('time to temp')
    
    tempc2 = tempc.tolist()
    tempc2.insert(0,"file name")
    tempc2.append('p/f')                                                        #add pass fail column
    tempc2 = np.array(tempc2)


    tempc3 = []                                                                 
    tempc2 = tempc2.tolist()                                                    
    for i in tempc2:                                            
        try:    
            tempc3.append(float(i))                                             
        except:     
            tempc3.append(i)                                                    #format list of temps         

    ws.append(tempc3)                                                           #add time-to data to excel
    timeTo3 = [ [] for _ in range(len(timeTo2.T))]
    for timeIndx in (range(len(timeTo2.T))):
        for u in timeTo2.T[timeIndx]:
            try:
                timeTo3[timeIndx].append(float(u))
            except:
                timeTo3[timeIndx].append(u)
    for time in timeTo3:
        ws.append(time)
        
    
    values = Reference(ws,min_col=1, min_row=2, max_col=len(tempc)+1, max_row=len(timeTo3)+1)  #grab data to put on chart         
    labels = Reference(ws,min_col=2, min_row=1, max_col=len(tempc)+1, max_row=1)

    chart = LineChart()

    chart.add_data(values,titles_from_data=True,from_rows=True)                                 #add data to line chart
    chart.set_categories(labels=labels)

    
    ws.add_chart(chart,'K5')

    wb.save(newFile)            



##########################              ANALYZE DATA FROM TC HEATSPREADER               ###################################
def analyzeTC():
    newFile = 'TC_Results.xlsx'                         #name of excel file to save data
    folder = 'dataTC'                                 #name of folder where TC raw data is held
    err = 0.2                                          #error criteria (temperatures must be at least 80% nominal)
    
    timeTo = []                                         #initialize lists to hold data 
    fullTemp = []
    fullTime = []

    tempc = np.arange(40,105,10)                        #temps to evaluate 

    print('File being used as "nominal":',os.listdir(folder)[0])

    for fileName in os.listdir(folder):                                            #loop through all files in folder
        file = open(''.join([''.join([folder,'/']),fileName]),'r')                 #open file
        filex = file.readlines()                                                   #read file line by line
        
        time = []                                                                   #initiate lists to parse out time and temp data
        temp = []
        absDif = []                                                                 
        goodsTemp = []
        goodsTime = []

    ##########  parse data out of file ################
        for lineIndx in range(len(filex)):                                                 #loop through each line in the file
            if 'TC-' in filex[lineIndx]:                                                   #check if line contains desired data
                filex[lineIndx] = filex[lineIndx].split()                                         #split line into individual values
                goodsTemp.append(float(filex[lineIndx][4].strip(',')))                     #add thermistor temp (c) to list
                goodsTime.append(float(filex[lineIndx][0][1:-1])/1000)                     #add time (sec) to list

        for TempIndx in range(len(goodsTemp)):                                             #loop through temp data 
            if TempIndx != 0:                                                              #skip the first data point (can be weird)
                if goodsTemp[TempIndx]-goodsTemp[TempIndx-1] >= 0.5 and goodsTemp[TempIndx] >=25:        #dont use data until ramp rate >.5 c/s and temp at least 25c 
                    time.append(goodsTime[TempIndx])
                    temp.append(goodsTemp[TempIndx])
                elif goodsTemp[TempIndx] >=50:                                             #ramp rate not a criteria after 50c
                    time.append(goodsTime[TempIndx])
                    temp.append(goodsTemp[TempIndx])

        time = np.array(time)-time[0]                                               #normalize time            

        for Temp in tempc:                                                             #loop through set of temps to analyze
            absDif = []                                                             
            indx = []
            if max(temp) >= Temp:                                                      #find temps that are closest to desired temps                                               
                for i in temp:                             
                    absDif.append(abs(Temp-i))
                indx = absDif.index(min(absDif))
                timeTo.append(time[indx])                                           #append time it takes to reach each desired temp
            else:
                timeTo.append(0)                                                    #if desired temp is not reached, add 0
        fullTemp.append(temp)
        fullTime.append(time)



    timeTo = [timeTo[i:i+len(tempc)] for i in range(0,len(timeTo),len(tempc))]      #format data
    timeTo = np.array(timeTo).T

    lens = []
    for timeIndx in range(len(fullTime)):                                                  #find file with longest runtime
        lens.append(len(fullTime[timeIndx]))
    longest = max(lens)


    sameLenTime = []
    sameLenTemp = []

    for timeIndx in range(len(fullTime)):                                                  #make everything the same length
        x = list(fullTime[timeIndx])
        y = list(fullTemp[timeIndx])
        while len(x) < longest:
            
            x.append(None)
            y.append(None)
        sameLenTemp.append(y)
        sameLenTime.append(x)

    timeTo2 = timeTo.tolist()
    timeTo2.insert(0,os.listdir(folder))                                             #insert file name to beginning of list
    timeTo2 = np.array(timeTo2,dtype=object)
  

    timesInterpTot = []                                                                         
    for i in range(len(fullTemp)):
        timesInterpTot.append(interp1d(fullTemp[i],fullTime[i]))                                     #interpolate time as function of temperature

    nominal = timesInterpTot[0]

    timesInterpNominal = []
    for TEMP in tempc:
        try:
            timesInterpNominal.append(nominal(TEMP))                                                           #list of times it takes to get to desired temps
        except:
            timesInterpNominal.append(0)                                                                    #if temp isn't reached, add 0
                                                                                                            #nominal list used as a baseline to compare all others

    pfTot = []
    for tempIndx in range(len(fullTemp)):
        temps = interp1d(fullTime[tempIndx],fullTemp[tempIndx])                                                   #interpolate temperature as function of time
        tempsInterp = []
        count = 0
        for times in timesInterpNominal:    
            try:
                tempsInterp.append(temps(times))                                                                #list temperature at teh time it takes nominal to reach desired temp
            except:
                try:
                    tempsInterp.append(temps(timesInterpTot[tempIndx](tempc[count])))                              
                except:
                    tempsInterp.append(0)                                                                   #if temp isn't reached, add 0
            count += 1


        count = 4
        pf = []
        for i in tempsInterp[4:]:                                                           
            if tempc[count]>i and tempc[count]-i > tempc[count]*(err):                                    #check if temp at specified time is at least 80% of nominal temp
                pf.append('f')                                                                              #looks at 80,90, and 100c
            else:
                pf.append('p')                                                                              #if at least 80% pass
            count += 1

        if 'f' not in pf:                                                                                   #if any temp fails in each run, the whole thing fails
            pfTot.append('pass')
        else:
            pfTot.append('fail')


    timeTo2 = timeTo2.tolist()                                                                              #format data
    timeTo2.append(np.array(pfTot))
    timeTo2 = np.array(timeTo2)


    toExcel(folder,longest,sameLenTime,sameLenTemp,newFile,tempc,timeTo2)                                                                                           #run function to create excel file
    print('Analysis Complete')





##########################              ANALYZE DATA FROM PCR HEATSPREADER               ###################################
def analyzePCR():
    newFile = 'PCR_Results.xlsx'                                                #name of excel file to upload data
    folder = 'DataPCR'                                                         #name of folder where raw data is held


    err = 0.05                                                            #error criteria (temperatures must be at least 95% nominal)
    timeTo = []
    fullTemp = []
    fullTime = []

    tempc = np.arange(40,105,10)
    print('File being used as "nominal":',os.listdir(folder)[0])
    for fileName in os.listdir(folder):                                            #loop through all files in "data" folder
        file = open(''.join([''.join([folder,'/']),fileName]),'r')                 #open file
        filex = file.readlines()                                                   #read file line by line
        time = []                                                                   #initiate lists to parse out time and temp data
        temp = []
        absDif = []                                                                 

        goodsTemp = []
        goodsTime = []

    ##########  parse .txt file ################
        for u in range(len(filex)):                                                 #loop through each line in the file
            if 'TC-' in filex[u]:                                                   #check if line contains desired data
                filex[u] = filex[u].split()                                         #split line into individual values
                goodsTemp.append(float(filex[u][4].strip(',')))                     #add thermistor temp (c) to list
                goodsTime.append(float(filex[u][0][1:-1])/1000)                     #add time (sec) to list


        for j in range(len(goodsTemp)):                                             #loop through temps
            if j != 0:                                                              
                if goodsTemp[j]-goodsTemp[j-1] >= 0.5 and goodsTemp[j] >=25:        #dont use data until ramp rate >.5 c/s and temp at least 25c 
                    time.append(goodsTime[j])
                    temp.append(goodsTemp[j])
                elif goodsTemp[j] >=50:                                             #ramp rate not a criteria after 50c
                    time.append(goodsTime[j])
                    temp.append(goodsTemp[j])

        time = np.array(time)-time[0]                                               #normalize time            

        
        for u in tempc:                                                             #loop through set of temps to analyze
            absDif = []                                                             
            indx = []
            if max(temp) >= u:                                                      #find temps that are closest to desired temps                                               
                for i in temp:                             
                    absDif.append(abs(u-i))
                indx = absDif.index(min(absDif))
                timeTo.append(time[indx])                                           #add time it takes to reach each desired temp
            else:
                timeTo.append(0)                                                    #if desired temp is not reached, add 0
        fullTemp.append(temp)
        fullTime.append(time)


    timeTo = [timeTo[i:i+len(tempc)] for i in range(0,len(timeTo),len(tempc))]      #format data
    timeTo = np.array(timeTo).T

    lens = []
    for i in range(len(fullTime)):                                                  #find file with longest runtime
        lens.append(len(fullTime[i]))
    longest = max(lens)


    sameLenTime = []
    sameLenTemp = []

    for i in range(len(fullTime)):                                                  #make everything the same length
        x = list(fullTime[i])
        y = list(fullTemp[i])
        while len(x) < longest:
            
            x.append(None)
            y.append(None)
        sameLenTemp.append(y)
        sameLenTime.append(x)

    timeTo2 = timeTo.tolist()
    timeTo2.insert(0,os.listdir(folder))                                             #insert file name to beginning of list
    timeTo2 = np.array(timeTo2,dtype=object)



    timesInterpTot = []                                                                         
    for i in range(len(fullTemp)):
        timesInterpTot.append(interp1d(fullTemp[i],fullTime[i]))                                     #interpolate time as function of temperature

    nominal = timesInterpTot[0]
    timesInterpNominal = []

    for i in tempc:
        try:
            timesInterpNominal.append(nominal(i))                                                           #list of times it takes to get to desired temps
        except:
            timesInterpNominal.append(0)                                                                    #if temp isn't reached, add 0
                                                                                                            #nominal list used as a baseline to compare all others

    pfTot = []
    for j in range(len(fullTemp)):
        temps = interp1d(fullTime[j],fullTemp[j])                                                   #interpolate temperature as function of time
        tempsInterp = []
        count = 0
        for i in timesInterpNominal:    
            try:
                tempsInterp.append(temps(i))                                                                #list temperature at teh time it takes nominal to reach desired temp
            except:
                try:
                    tempsInterp.append(temps(timesInterpTot[j](tempc[count])))                              
                except:
                    tempsInterp.append(0)                                                                   #if temp isn't reached, add 0
            count += 1




        count = 4
        pf = []

        for i in tempsInterp[4:]:                                                           
            if tempc[count]>i and tempc[count]-i > tempc[count]*(err):                                    #check if temp at specified time is at least 95% nominal
                pf.append('f')                                                                              
            else:
                pf.append('p')                                                                              #if at least 95% pass
            count += 1

        if 'f' not in pf:                                                                                   #if any temp fails in each run, the whole thing fails
            pfTot.append('pass')
        else:
            pfTot.append('fail')


    timeTo2 = timeTo2.tolist()                                                                              #format data
    timeTo2.append(np.array(pfTot))
    timeTo2 = np.array(timeTo2)



    toExcel(folder,longest,sameLenTime,sameLenTemp,newFile,tempc,timeTo2)                                   #run function to create excel file
    print('Analysis Complete')




#############################               STANDALONE EXE                          #########################
root = Tk()                                                                                             #initialize user interface window
root.geometry("400x400")                                                                                #set size of window

root.title("Heat Spreader Characterization Analysis Tool")                                                             #name ui window

run = Button(root,height=2,width=20,text='Analyze PCR',command=lambda:analyzePCR())                            #create button that runs the above code when pushed
run2 = Button(root,height=2,width=20,text='Analyze TC',command=lambda:analyzeTC())                            #create button that runs the above code when pushed
run.pack()
run2.pack()
mainloop()








                                                                          


