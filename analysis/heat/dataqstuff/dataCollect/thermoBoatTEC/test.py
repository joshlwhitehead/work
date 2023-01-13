import os
import numpy as np
import matplotlib.pyplot as plt
import pandas as pd
import openpyxl as op
import xlsxwriter
from scipy import stats
from scipy import interpolate as interp

from openpyxl.chart import LineChart, Reference, Series
from statsmodels.stats.anova import anova_lm
from statsmodels.formula.api import ols
from statsmodels.stats.multicomp import pairwise_tukeyhsd

newFile = 'TC.xlsx'
folder = 'dataTC'
timeTo = []
fullTemp = []
fullTime = []
fullDerivTemp = []
tempc = np.arange(40,115,10)
# print(len(os.listdir(folder)))
for k in os.listdir(folder):
    fileName = k#'Thermalboat 20221207 Rebuilt Run 3.txt'
    file = open(''.join([''.join([folder,'/']),fileName]),'r')
    filex = file.readlines()
    # tempC = 90
    time = []
    temp = []
    absDif = []

    goodsTemp = []
    goodsTime = []
  
    for u in range(len(filex)):
        if 'TC-' in filex[u]:
            filex[u] = filex[u].split()
            goodsTemp.append(float(filex[u][4][:-1]))
            goodsTime.append(float(filex[u][0][1:-1])/1000)


    
    for j in range(len(goodsTemp)):
       
        if j != 0:
            if goodsTemp[j]-goodsTemp[j-1] >= 0.5 and goodsTemp[j] >=25:
                
                time.append(goodsTime[j])
                temp.append(goodsTemp[j])
            elif goodsTemp[j] >=50:
                time.append(goodsTime[j])
                temp.append(goodsTemp[j])
    
    time = np.array(time)-time[0]


    
    for u in tempc:
        absDif = []
        indx = []
        if max(temp) >= u:
            for i in temp:
         
                absDif.append(abs(u-i))

            indx = absDif.index(min(absDif))
         
            
            timeTo.append(time[indx])
        else:
            timeTo.append(0)
    fullTemp.append(temp)
    fullTime.append(time)

    timeTemp = np.array([time,temp]).T

    a,b,c = np.polyfit(time,temp,2)
    fullDerivTemp.append(2*a*time+b)


short = 9999
for i in fullTemp:
    if len(i) < short:
        short = len(i)

dFTest = pd.DataFrame({'time':fullTime[:short],'temp':fullTemp[:short]})

fileUnpack = []
tempUnpack = []
count = 0
for i in fullTemp[:short]:
    for u in i:
        fileUnpack.append(os.listdir(folder)[count])
        tempUnpack.append(u)
    count +=1


dfTestAvg = pd.DataFrame({'file':fileUnpack,'temp':tempUnpack})
# dfTestAvg.boxplot('temp',by='file')
# plt.show()



timeTo = [timeTo[i:i+len(tempc)] for i in range(0,len(timeTo),len(tempc))]
timeTo = np.array(timeTo).T

lens = []
for i in range(len(fullTime)):
    lens.append(len(fullTime[i]))
longest = max(lens)
longestTime = fullTime[lens.index(longest)]

sameLenTime = []
sameLenTemp = []
sameLenDeriv = []

for i in range(len(fullTime)):
    x = list(fullTime[i])
    y = list(fullTemp[i])
    z = list(fullDerivTemp[i])
    while len(x) < longest:
        
        x.append(None)
        y.append(None)
        z.append(None)
    sameLenTemp.append(y)
    sameLenTime.append(x)
    sameLenDeriv.append(z)

timeTo2 = timeTo.tolist()

timeTo2.insert(0,os.listdir(folder))

timeTo2 = np.array(timeTo2,dtype=object)


def toExcel():
    it = np.arange(0,len(os.listdir(folder)))



    fullDict = {}
    for i in range(len(os.listdir(folder))):
        x = [os.listdir(folder)[i]]
    
        while len(x) < longest:
            x.append(None)
        fullDict[''.join(['file name ',str(it[i])])] = x
        fullDict[''.join(['normalized time (sec) ',str(it[i])])] = sameLenTime[i]
        fullDict[''.join(['temp (c) ',str(it[i])])] = sameLenTemp[i]


    dFTot = pd.DataFrame(fullDict)
    writer = pd.ExcelWriter(newFile,engine='xlsxwriter')
    dFTot.to_excel(writer,sheet_name='full')
    wb = writer.book
    ws = writer.sheets['full']
    chart = wb.add_chart({'type':'line'})

    count = 2
    for i in range(len(sameLenTemp)):
        chart.add_series({
            'categories':['full',1,count,len(sameLenTemp[0]),count],#['full',1,lens.index(longest)*3+2,len(sameLenTemp[lens.index(longest)]),lens.index(longest)*3+2],
            'values':['full',1,count+1,len(sameLenTemp[0]),count+1],
            'name':['full',1,count-1]
            })
        count += 3
    
 
    chart.set_x_axis({'name':'Time (sec)'})
    chart.set_y_axis({'name':'Temp (c)'})
    

    ws.insert_chart('D2',chart)
    writer.save()





    



    wb = op.load_workbook(newFile)
    ws = wb.create_sheet('time to temp')
    
    tempc2 = tempc.tolist()
    tempc2.insert(0,"file name")
    tempc2.append('p/f')
    tempc2 = np.array(tempc2)
  
    timeToDict = {}
    for i in range(len(tempc2)):
        timeToDict[tempc2[i]] = list(timeTo2[i])
    dF = pd.DataFrame(timeTo2.T,columns=tempc2)
    tempc3 = []
    tempc2 = tempc2.tolist()
    for i in tempc2:
        try:
            tempc3.append(float(i))
        except:
            tempc3.append(i)

    ws.append(tempc3)
    timeTo3 = [ [] for _ in range(len(timeTo2.T))]
    for i in (range(len(timeTo2.T))):
        
        for u in timeTo2.T[i]:
            
            try:
                timeTo3[i].append(float(u))
            except:
                timeTo3[i].append(u)
    for i in timeTo3:

        ws.append(i)
        
    
    values = Reference(ws,min_col=1, min_row=2, max_col=len(tempc)+1, max_row=len(timeTo3)+1)
    labels = Reference(ws,min_col=2, min_row=1, max_col=len(tempc)+1, max_row=1)
  
    chart = LineChart()

    chart.add_data(values,titles_from_data=True,from_rows=True)
    chart.set_categories(labels=labels)
   
    
    ws.add_chart(chart,'K5')
  
   
    wb.save(newFile)

    
        





timeToComp = []
for i in timeTo.T:
    timeToComp.append(i-timeTo.T[0])


crit = 0.05
count2 = 0
pfTot = []
for i in timeToComp:
    count = 0
    pf = []
    for u in i:
        
        if u < crit*tempc[count]:
            pf.append(1)
        else:
            pf.append(0)
        count +=1
    if 0 not in pf:
 
        pfTot.append('pass')
    else:
        pfTot.append('fail')
    count2 += 1

timeTo2 = timeTo2.tolist()
timeTo2.append(np.array(pfTot))
timeTo2 = np.array(timeTo2)

print(timeTo[0])
print(tempc)





# toExcel()
# count = 0
# for i in timeTo.T:

#     if count == 0:
#         plt.plot(tempc,timeToComp[count],'k',lw=5,label='nominal')
#     else:
#         plt.plot(tempc,timeToComp[count])
#     count += 1
# plt.grid()
# plt.xlabel('Temp (c)')
# plt.ylabel('Time to Reach Temp (sec)')
# plt.title('Time to Temp Compared to Nominal')
# plt.legend()



# plt.figure()
# count = 0
# for i in timeTo.T:
#     if count == 0:
#         plt.plot(tempc,i,'k',lw=5,label='nominal')
#     else:
#         plt.plot(tempc,i)
#     count += 1
# plt.legend()
# plt.grid() 
# plt.xlabel('Temp (c)')
# plt.ylabel('Time to Reach Temp (sec)')
# plt.title('Time to Temp')




def r2(y,fit):
    st = sum((y-np.average(y))**2)
    sr = sum((y-fit)**2)
    r2 = 1-sr/st
    return r2

count = 0

for i in fullTemp[:]:
    plt.plot(fullTime[count][:short],i[:short],label=pfTot[count])
    # print(r2(np.array(fullTemp[0][:short]),np.array(i[:short])))
    count += 1
plt.legend()
plt.grid()
plt.show()

