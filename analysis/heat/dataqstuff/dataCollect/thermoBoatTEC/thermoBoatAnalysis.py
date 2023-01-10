import os
import numpy as np
import matplotlib.pyplot as plt
import pandas as pd
import openpyxl as op
import xlsxwriter
from scipy import stats
import statsmodels.api as sm

newFile = 'PCR.xlsx'
folder = 'data'
timeTo = []
fullTemp = []
fullTime = []
fullDerivTemp = []
tempc = np.arange(40,115,10)
# print(len(os.listdir(folder)))
for k in os.listdir(folder):
    fileName = k#'Thermalboat 20221207 Rebuilt Run 3.txt'
    file = open(''.join([''.join([folder,'/']),fileName]),'r')
    filex = file.readlines()
    # tempC = 90
    time = []
    temp = []
    absDif = []

    goodsTemp = []
    goodsTime = []
  
    for u in range(len(filex)):
        if 'TC-' in filex[u]:
            filex[u] = filex[u].split()
            goodsTemp.append(float(filex[u][4][:-1]))
            goodsTime.append(float(filex[u][0][1:-1])/1000)


            # if float(goodsTemp[count+1])-float(goodsTemp[count])>=0.5 and float(filex[count][4][:-1]) >= 25:
                
            #     time.append(float(filex[count][0][1:-1])/1000)
            #     temp.append(float(filex[count][4][:-1]))
            #     # absDif.append(abs(tempC-float(u[5][:-1])))
    # print(goodsTemp)
    
    for j in range(len(goodsTemp)):
       
        if j != 0:
            if goodsTemp[j]-goodsTemp[j-1] >= 0.5 and goodsTemp[j] >=25:
                
                time.append(goodsTime[j])
                temp.append(goodsTemp[j])
            elif goodsTemp[j] >=50:
                time.append(goodsTime[j])
                temp.append(goodsTemp[j])
    
    time = np.array(time)-time[0]


    
    for u in tempc:
        absDif = []
        indx = []
        if max(temp) >= u:
            for i in temp:
                # if abs(u-i) <= 0.5
                absDif.append(abs(u-i))

            indx = absDif.index(min(absDif))
            # print(time[indx])
            
            timeTo.append(time[indx])
        else:
            timeTo.append(0)
    fullTemp.append(temp)
    fullTime.append(time)

    timeTemp = np.array([time,temp]).T
    # wb = op.load_workbook(newFile)
    a,b,c = np.polyfit(time,temp,2)
    fullDerivTemp.append(2*a*time+b)
#     plt.plot(time,temp,label=fileName[:7])
# plt.grid()   
# plt.legend()
# plt.show()
# fullTemp = np.array(fullTemp)
# fullTime = np.array(fullTime)
# print(file.read())
dFTest = pd.DataFrame({'time':fullTime[0],'temp':fullTemp[0]})

formula = 'temp ~ time'
model = sm.formula.ols(formula,dFTest)
model_fitted = model.fit()
# print(model_fitted.summary())
    
timeTo = [timeTo[i:i+len(tempc)] for i in range(0,len(timeTo),len(tempc))]
timeTo = np.array(timeTo).T

lens = []
for i in range(len(fullTime)):
    lens.append(len(fullTime[i]))

sameLenTime = []
sameLenTemp = []
sameLenDeriv = []
longest = max(lens)
for i in range(len(fullTime)):
    x = list(fullTime[i])
    y = list(fullTemp[i])
    z = list(fullDerivTemp[i])
    while len(x) < longest:
        
        x.append(None)
        y.append(None)
        z.append(None)
    sameLenTemp.append(y)
    sameLenTime.append(x)
    sameLenDeriv.append(z)

timeTo2 = timeTo.tolist()

timeTo2.insert(0,os.listdir(folder))
# print(type(timeTo2))
timeTo2 = np.array(timeTo2,dtype=object)


def toExcel():
    it = np.arange(0,len(os.listdir(folder)))
    # print(it)


    fullDict = {}
    for i in range(len(os.listdir(folder))):
        x = [os.listdir(folder)[i]]
        # print(x)
        while len(x) < longest:
            x.append(None)
        fullDict[''.join(['file name ',str(it[i])])] = x
        fullDict[''.join(['normalized time (sec) ',str(it[i])])] = sameLenTime[i]
        fullDict[''.join(['temp (c) ',str(it[i])])] = sameLenTemp[i]


    dFTot = pd.DataFrame(fullDict)
    writer = pd.ExcelWriter(newFile,engine='xlsxwriter')
    dFTot.to_excel(writer,sheet_name='full')
    wb = writer.book
    ws = writer.sheets['full']
    chart = wb.add_chart({'type':'line'})

    count = 2
    for i in range(len(sameLenTemp)):
        chart.add_series({
            'categories':['full',1,count,len(sameLenTemp[0]),count],
            'values':['full',1,count+1,len(sameLenTemp[0]),count+1],
            'name':['full',1,count-1]
            })
        count += 3
    chart.set_x_axis({'name':'Time (sec)'})
    chart.set_y_axis({'name':'Temp (c)'})

    ws.insert_chart('D2',chart)
    writer.save()









    wb = op.load_workbook(newFile)
    tempc2 = tempc.tolist()
    tempc2.insert(0,'file name')
    tempc2.append('p/f')
    tempc2 = np.array(tempc2)
  
    with pd.ExcelWriter(newFile,engine='openpyxl') as writer:
        writer.book = wb
        writer.sheets = {worksheet.title:worksheet for worksheet in wb.worksheets}
        dF = pd.DataFrame(timeTo2.T,columns=tempc2)
        dF.to_excel(writer,'time to temp')
        # wb = writer.book
        # ws = writer.sheets['time to temp']
        # chart = wb.add_chart({'type':'line'})

        # count = 1
        # for i in range(len(tempc2)-1):
        #     chart.add_series({
        #         'categories':['time to temp',0,2,0,len(tempc2)-1],
        #         'values':['time to temp',count,2,count,len(tempc2)-1],
        #         'name':['time to temp',count,1]
        #         })
        #     count += 1
        # chart.set_x_axis({'name':'Temp (c)'})
        # chart.set_y_axis({'name':'Time (sec)'})

        # ws.insert_chart('J2',chart)
        writer.save()


    fullDict = {}
    for i in range(len(os.listdir(folder))):
        x = [os.listdir(folder)[i]]
        while len(x) < longest:
            x.append(None)
        fullDict[''.join(['file name ',str(it[i])])] = x
        fullDict[''.join(['normalized time (sec) ',str(it[i])])] = sameLenTime[i]
        fullDict[''.join(['dTdt ',str(it[i])])] = sameLenDeriv[i]


    with pd.ExcelWriter(newFile,engine='openpyxl') as writer:
        writer.book = wb
        writer.sheets = {worksheet.title:worksheet for worksheet in wb.worksheets}
        dF = pd.DataFrame(fullDict)
        dF.to_excel(writer,'deriv')
        writer.save()

    # with pd.ExcelWriter(newFile,engine='xlsxwriter') as writer:
    
        





timeToComp = []
for i in timeTo.T:
    timeToComp.append(i-timeTo.T[0])



crit = 0.05
count2 = 0
pfTot = []
for i in timeToComp:
    count = 0
    pf = []
    for u in i:
        if u < crit*tempc[count]:
            pf.append(1)
        else:
            pf.append(0)
        count +=1
    if 0 not in pf:
        print(timeTo2[0][count2])
        pfTot.append('pass')
    else:
        pfTot.append('fail')
    count2 += 1
timeTo2 = timeTo2.tolist()
timeTo2.append(np.array(pfTot))
timeTo2 = np.array(timeTo2)
# print(timeTo2)





toExcel()
count = 0
for i in timeTo.T:
    # plt.plot(tempc,i)
    # print(np.average(timeToComp[count]))
    if count == 0:
        plt.plot(tempc,timeToComp[count],'k',lw=5,label='nominal')
    else:
        plt.plot(tempc,timeToComp[count])
    count += 1
plt.grid()
plt.xlabel('Temp (c)')
plt.ylabel('Time to Reach Temp (sec)')
plt.title('Time to Temp Compared to Nominal')
plt.legend()



plt.figure()
count = 0
for i in timeTo.T:
    if count == 0:
        plt.plot(tempc,i,'k',lw=5,label='nominal')
    else:
        plt.plot(tempc,i)
    count += 1
plt.legend()
plt.grid() 
plt.xlabel('Temp (c)')
plt.ylabel('Time to Reach Temp (sec)')
plt.title('Time to Temp')
# plt.show()